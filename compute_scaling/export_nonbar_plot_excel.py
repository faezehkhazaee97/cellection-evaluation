#!/usr/bin/env python3
"""The numbers behind the line and scatter plots of covid_scaling.py, as an Excel workbook.

Run after covid_scaling.py:
  python export_nonbar_plot_excel.py [--results-dir D] [--output F]

Reads scaling_runtime_gpu.csv from --results-dir (default $RESULTS_DIR/R1-Q13/Code/Results/Covid) and
writes --output (default <results-dir>/covid_nonbar_plot_values.xlsx) with one sheet per line or
scatter plot; the three bar plots are left out. Each sheet opens with the plot definition (title,
file, source CSV, axes, hue, point size). Line-plot sheets then give the mean, std and count over the
seed x fold runs at each (train_patients, cells_per_patient) point, followed by the raw rows.
Scatter-plot sheets give every run and the numpy.polyfit degree-1 line the figure overlays (slope,
intercept, fitted values).
"""

# Paths come from environment variables (see experiments-public/README.md).
import os
RESULTS_DIR = os.environ.get("RESULTS_DIR", "results")  # output root

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Excel limits sheet names to 31 characters.
MAX_SHEET_NAME_LEN = 31
# Column width: longest cell text + COLUMN_PADDING, capped at MAX_COLUMN_WIDTH.
COLUMN_PADDING = 2
MAX_COLUMN_WIDTH = 38

LINE_PLOTS = [
    {
        "sheet": "Epoch Time vs Cells",
        "title": "Average Epoch Time vs. Cells per Patient",
        "plot_file": "epoch_time_vs_cells.pdf",
        "x": "cells_per_patient",
        "y": "epoch_time_sec",
        "hue": "train_patients",
        "x_label": "Cells per Patient",
        "y_label": "Time per Epoch (sec)",
    },
    {
        "sheet": "Total Time vs Cells",
        "title": "Total Training Time vs. Cells per Patient",
        "plot_file": "total_time_vs_cells.pdf",
        "x": "cells_per_patient",
        "y": "total_time_sec",
        "hue": "train_patients",
        "x_label": "Cells per Patient",
        "y_label": "Total Time (sec)",
    },
    {
        "sheet": "GPU Memory vs Cells",
        "title": "GPU Memory Usage vs. Cells per Patient",
        "plot_file": "gpu_memory_vs_cells.pdf",
        "x": "cells_per_patient",
        "y": "gpu_memory_mb",
        "hue": "train_patients",
        "x_label": "Cells per Patient",
        "y_label": "Peak GPU Memory (MB)",
    },
]

SCATTER_PLOTS = [
    {
        "sheet": "GPU Mem vs Total Cells",
        "title": "GPU Memory vs Total Cells in Training",
        "plot_file": "gpu_memory_vs_total_cells.pdf",
        "x": "total_cells_train",
        "y": "gpu_memory_mb",
        "hue": "train_patients",
        "size": "cells_per_patient",
        "x_label": "Total Cells (Train)",
        "y_label": "GPU Peak Memory (MB)",
    },
    {
        "sheet": "Time vs Total Cells",
        "title": "Total Training Time vs Total Cells in Training",
        "plot_file": "total_time_vs_total_cells.pdf",
        "x": "total_cells_train",
        "y": "total_time_sec",
        "hue": "train_patients",
        "size": "cells_per_patient",
        "x_label": "Total Cells (Train)",
        "y_label": "Total Training Time (sec)",
    },
]


def default_results_dir() -> Path:
    return Path(RESULTS_DIR) / "R1-Q13" / "Code" / "Results" / "Covid"


def line_plot_values(runtime_df: pd.DataFrame, plot: dict) -> pd.DataFrame:
    """Mean, std and run count of plot["y"] at each (hue, x) point."""
    grouped = runtime_df.groupby([plot["hue"], plot["x"]], sort=True)[plot["y"]]
    out = grouped.agg(["mean", "std", "count"]).reset_index()
    out = out.rename(
        columns={
            plot["hue"]: "train_patients",
            plot["x"]: "cells_per_patient",
            "mean": f"{plot['y']}_mean",
            "std": f"{plot['y']}_std",
            "count": "n_rows",
        }
    )
    return out


def scatter_plot_points(runtime_df: pd.DataFrame, plot: dict) -> pd.DataFrame:
    """One row per run with the plotted columns plus epochs_run and n_genes."""
    columns = [
        "seed",
        "fold",
        plot["hue"],
        plot["size"],
        plot["x"],
        plot["y"],
        "epochs_run",
        "n_genes",
    ]
    return runtime_df[columns].copy()


def best_fit_values(runtime_df: pd.DataFrame, plot: dict) -> pd.DataFrame:
    """The add_best_fit_line() fit from covid_scaling.py, evaluated at the sorted finite x values.

    Returns an empty table when fewer than two finite points exist.
    """
    x = runtime_df[plot["x"]].to_numpy(dtype=float)
    y = runtime_df[plot["y"]].to_numpy(dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    if mask.sum() < 2:
        return pd.DataFrame(columns=[plot["x"], f"{plot['y']}_fit", "slope", "intercept"])
    slope, intercept = np.polyfit(x[mask], y[mask], 1)
    x_sorted = np.sort(x[mask])
    return pd.DataFrame(
        {
            plot["x"]: x_sorted,
            f"{plot['y']}_fit": slope * x_sorted + intercept,
            "slope": slope,
            "intercept": intercept,
        }
    )


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(
            max_length + COLUMN_PADDING, MAX_COLUMN_WIDTH
        )


def style_sheet(worksheet, title_rows: list[int], header_rows: list[int]) -> None:
    """Colour the title and header rows (1-based), freeze the top two rows and size the columns."""
    title_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_fill = PatternFill(fill_type="solid", fgColor="EDEDED")

    for row_idx in title_rows:
        for cell in worksheet[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = title_fill

    for row_idx in header_rows:
        for cell in worksheet[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

    worksheet.freeze_panes = "A3"
    autosize_columns(worksheet)


def write_table(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    table: pd.DataFrame,
    startrow: int,
) -> tuple[int, int, int]:
    """Write a one-cell title row and `table` below it, starting at 0-based row `startrow`.

    Returns (next_start, title_row, header_row): the 0-based start of the next table, two blank rows
    further down, and the 1-based Excel rows of this title and header, for styling.
    """
    pd.DataFrame([[title]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=startrow
    )
    table.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow + 1)
    title_row = startrow + 1
    header_row = startrow + 2
    next_start = startrow + len(table) + 4
    return next_start, title_row, header_row


def plot_metadata(plot: dict, source_csv: str) -> pd.DataFrame:
    rows = [
        ["Plot title", plot["title"]],
        ["Plot file", plot["plot_file"]],
        ["Source CSV", source_csv],
        ["X axis", plot["x_label"]],
        ["Y axis", plot["y_label"]],
    ]
    if "hue" in plot:
        rows.append(["Hue", plot["hue"]])
    if "size" in plot:
        rows.append(["Point size", plot["size"]])
    return pd.DataFrame(rows, columns=["field", "value"])


def write_line_sheet(writer: pd.ExcelWriter, runtime_df: pd.DataFrame, plot: dict) -> None:
    """Sheet layout for a line plot: plot definition, per-point summary, raw rows."""
    sheet_name = plot["sheet"][:MAX_SHEET_NAME_LEN]
    row = 0
    title_rows = []
    header_rows = []

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Source files and plot definition",
        plot_metadata(plot, "scaling_runtime_gpu.csv"),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    values = line_plot_values(runtime_df, plot)
    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Plot line values: mean across seed/fold runs for each line point",
        values,
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    raw_columns = [
        "seed",
        "fold",
        "train_patients",
        "cells_per_patient",
        "total_cells_train",
        plot["y"],
        "epochs_run",
        "n_genes",
    ]
    raw = runtime_df[raw_columns].copy()
    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Raw rows used by seaborn lineplot",
        raw,
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    style_sheet(writer.sheets[sheet_name], title_rows, header_rows)


def write_scatter_sheet(writer: pd.ExcelWriter, runtime_df: pd.DataFrame, plot: dict) -> None:
    """Sheet layout for a scatter plot: plot definition, every point, least-squares line."""
    sheet_name = plot["sheet"][:MAX_SHEET_NAME_LEN]
    row = 0
    title_rows = []
    header_rows = []

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Source files and plot definition",
        plot_metadata(plot, "scaling_runtime_gpu.csv"),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    points = scatter_plot_points(runtime_df, plot)
    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Raw scatter points used by the plot",
        points,
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    fit = best_fit_values(runtime_df, plot)
    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Best-fit line values from numpy.polyfit",
        fit,
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    style_sheet(writer.sheets[sheet_name], title_rows, header_rows)


def export_workbook(results_dir: Path, output_path: Path) -> None:
    runtime_path = results_dir / "scaling_runtime_gpu.csv"
    if not runtime_path.is_file():
        raise FileNotFoundError(runtime_path)
    runtime_df = pd.read_csv(runtime_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for plot in LINE_PLOTS:
            write_line_sheet(writer, runtime_df, plot)
        for plot in SCATTER_PLOTS:
            write_scatter_sheet(writer, runtime_df, plot)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create Excel workbook for R1-Q13 non-bar scaling plots."
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=default_results_dir(),
        help="Folder containing scaling_runtime_gpu.csv and plot files.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to <results-dir>/covid_nonbar_plot_values.xlsx.",
    )
    args = parser.parse_args()

    output_path = args.output or args.results_dir / "covid_nonbar_plot_values.xlsx"
    export_workbook(args.results_dir, output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
