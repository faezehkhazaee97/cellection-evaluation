#!/usr/bin/env python3
"""One Excel workbook holding every number behind the cell-count before/after plots, a sheet per task.

Run it after all training and plot scripts in cell_count_imbalance/:
  python export_grouped_bin_excel.py [--results-dir D] [--output F]

For each task in DATASETS (Cardio, Cardio biosample, Covid Ziegler, Lupus disease, Lupus race) it reads
the before and after per-bin CSVs and per-(seed, fold) CSVs from --results-dir (default
$RESULTS_DIR/R1-q8/Results). Each sheet stacks four tables: the source files (task label, the two plot
PDFs, the four CSVs); the per-bin values drawn in the grouped bar PDFs, ordered by the bin's numeric
left edge; a per-condition summary (folds averaged within each seed, then the mean and sample std
across seeds); and the raw per-(seed, fold) rows. The workbook goes to --output, by default
<results-dir>/grouped_bin_plot_values.xlsx.
"""

# Paths come from environment variables (see experiments-public/README.md).
import os
RESULTS_DIR = os.environ.get("RESULTS_DIR", "results")  # output root

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Excel limits sheet names to 31 characters.
MAX_SHEET_NAME_LEN = 31
# Column width: longest cell text + COLUMN_PADDING, capped at MAX_COLUMN_WIDTH.
COLUMN_PADDING = 2
MAX_COLUMN_WIDTH = 42
TITLE_FILL_COLOR = "D9EAF7"
HEADER_FILL_COLOR = "EDEDED"

DATASETS = [
    {
        "sheet": "cardio",
        "label": "Cardio",
        "before_bins": "cardio/CELECTION/4-bin/per_bin_metrics.csv",
        "after_bins": "cardio/CELECTION/4-bin_downsample/per_bin_metrics.csv",
        "before_raw": "cardio/CELECTION/4-bin/results_summary.csv",
        "after_raw": "cardio/CELECTION/4-bin_downsample/summary_downsampled.csv",
        "auroc_plot": "cardio/CELECTION/4-bin_plots/cardio_4bin_auroc_before_after.pdf",
        "macrof1_plot": "cardio/CELECTION/4-bin_plots/cardio_4bin_macrof1_before_after.pdf",
    },
    {
        "sheet": "cardio_biosample",
        "label": "Cardio Biosample",
        "before_bins": "cardio_biosample/CELECTION/4-bin/per_bin_metrics.csv",
        "after_bins": "cardio_biosample/CELECTION/4-bin_downsample/per_bin_metrics.csv",
        "before_raw": "cardio_biosample/CELECTION/4-bin/results_summary.csv",
        "after_raw": "cardio_biosample/CELECTION/4-bin_downsample/summary_downsampled.csv",
        "auroc_plot": "cardio_biosample/CELECTION/4-bin_plots/cardiobiosample_4bin_auroc_before_after.pdf",
        "macrof1_plot": "cardio_biosample/CELECTION/4-bin_plots/cardiobiosample_4bin_macrof1_before_after.pdf",
    },
    {
        "sheet": "covid_Ziegler",
        "label": "Covid Ziegler",
        "before_bins": "covid_Ziegler/CELECTION/bin_8/cellection_global_bins.csv",
        "after_bins": "covid_Ziegler/CELECTION/bin_8_subsample_retrain/cellection_global_bins_retrain_downsampled.csv",
        "before_raw": "covid_Ziegler/CELECTION/bin_8/cellection_results_summary.csv",
        "after_raw": "covid_Ziegler/CELECTION/bin_8_subsample_retrain/cellection_results_summary.csv",
        "auroc_plot": "covid_Ziegler/CELECTION/auroc_grouped_bins_6.pdf",
        "macrof1_plot": "covid_Ziegler/CELECTION/macrof1_grouped_bins_6.pdf",
    },
    {
        "sheet": "lupus_disease",
        "label": "Lupus Disease",
        "before_bins": "lupus_disease/CELECTION/4-bin/per_bin_metrics.csv",
        "after_bins": "lupus_disease/CELECTION/4-bin_downsample/per_bin_metrics.csv",
        "before_raw": "lupus_disease/CELECTION/4-bin/results_summary.csv",
        "after_raw": "lupus_disease/CELECTION/4-bin_downsample/summary_downsampled.csv",
        "auroc_plot": "lupus_disease/CELECTION/4-bin_plots/lupusdisease_4bin_auroc_before_after.pdf",
        "macrof1_plot": "lupus_disease/CELECTION/4-bin_plots/lupusdisease_4bin_macrof1_before_after.pdf",
    },
    {
        "sheet": "lupus_race",
        "label": "Lupus Race",
        "before_bins": "lupus_race/CELECTION/4-bin/per_bin_metrics.csv",
        "after_bins": "lupus_race/CELECTION/4-bin_downsample/per_bin_metrics.csv",
        "before_raw": "lupus_race/CELECTION/4-bin/results_summary.csv",
        "after_raw": "lupus_race/CELECTION/4-bin_downsample/summary_downsampled.csv",
        "auroc_plot": "lupus_race/CELECTION/4-bin_plots/lupusrace_4bin_auroc_before_after.pdf",
        "macrof1_plot": "lupus_race/CELECTION/4-bin_plots/lupusrace_4bin_macrof1_before_after.pdf",
    },
]

METRICS = ["AUROC", "MacroF1"]
CONDITIONS = [("Before Downsampling", "before"), ("After Downsampling", "after")]


def default_results_dir() -> Path:
    return Path(RESULTS_DIR) / "R1-q8" / "Results"


def left_edge(interval_str: str) -> float:
    clean = (
        str(interval_str)
        .replace("(", "")
        .replace("[", "")
        .replace("]", "")
        .replace(")", "")
    )
    return float(clean.split(",")[0])


def sort_by_cell_bin(df: pd.DataFrame) -> pd.DataFrame:
    """Order rows by the bin's numeric left edge (text order would misplace bins), then by condition."""
    ordered_bins = sorted(df["cell_bin"].dropna().unique(), key=left_edge)
    out = df.copy()
    out["cell_bin"] = pd.Categorical(out["cell_bin"], categories=ordered_bins, ordered=True)
    return out.sort_values(["cell_bin", "condition"], kind="stable").reset_index(drop=True)


def load_plot_values(results_dir: Path, dataset: dict) -> pd.DataFrame:
    """Before and after per-bin tables of one task, stacked and tagged by condition."""
    before = pd.read_csv(results_dir / dataset["before_bins"])
    after = pd.read_csv(results_dir / dataset["after_bins"])
    before["condition"] = "Before Downsampling"
    after["condition"] = "After Downsampling"
    before["plot_source"] = "AUROC and MacroF1 grouped bin PDFs"
    after["plot_source"] = "AUROC and MacroF1 grouped bin PDFs"
    combined = pd.concat([before, after], ignore_index=True)
    columns = ["plot_source", "condition", "cell_bin", "AUROC", "MacroF1", "n_donors"]
    return sort_by_cell_bin(combined[columns])


def load_seed_fold_values(results_dir: Path, dataset: dict) -> pd.DataFrame:
    """Before and after per-(seed, fold) tables of one task, stacked and tagged by condition."""
    frames = []
    for condition, prefix in CONDITIONS:
        df = pd.read_csv(results_dir / dataset[f"{prefix}_raw"])
        df.insert(0, "condition", condition)
        frames.append(df)
    return pd.concat(frames, ignore_index=True)


def summarize_seed_fold(raw_df: pd.DataFrame) -> pd.DataFrame:
    """Per condition, average the folds within each seed, then take the mean and std across seeds.

    std is the sample std (ddof=1), and 0.0 when there is only one seed.
    """
    rows = []
    grouped = raw_df.groupby("condition", sort=False)
    for condition, df in grouped:
        per_seed = df.groupby("seed", sort=True)[METRICS].mean()
        row = {
            "condition": condition,
            "n_seeds": df["seed"].nunique(),
            "n_folds": df["fold"].nunique(),
            "n_rows": len(df),
        }
        for metric in METRICS:
            row[f"{metric}_mean"] = per_seed[metric].mean()
            row[f"{metric}_std"] = per_seed[metric].std(ddof=1) if len(per_seed) > 1 else 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(
            max_length + COLUMN_PADDING, MAX_COLUMN_WIDTH
        )


def style_title_row(worksheet, row_idx: int, fill_color: str) -> None:
    fill = PatternFill(fill_type="solid", fgColor=fill_color)
    for cell in worksheet[row_idx]:
        cell.font = Font(bold=True)
        cell.fill = fill


def style_sheet(worksheet, title_rows: list[int], header_rows: list[int]) -> None:
    """Colour the title and header rows (1-based) and freeze the top two rows."""
    for row_idx in title_rows:
        style_title_row(worksheet, row_idx, TITLE_FILL_COLOR)
    for row_idx in header_rows:
        style_title_row(worksheet, row_idx, HEADER_FILL_COLOR)
    worksheet.freeze_panes = "A3"
    autosize_columns(worksheet)


def write_dataset_sheet(
    writer: pd.ExcelWriter,
    sheet_name: str,
    plot_values: pd.DataFrame,
    seed_summary: pd.DataFrame,
    raw_values: pd.DataFrame,
    dataset: dict,
) -> None:
    """Write the four tables of one task one below another, each under a one-cell title row.

    pandas `startrow` is 0-based and openpyxl rows are 1-based, hence the `row + 1` bookkeeping.
    """
    metadata = pd.DataFrame(
        [
            ["Dataset", dataset["label"]],
            ["AUROC plot", dataset["auroc_plot"]],
            ["MacroF1 plot", dataset["macrof1_plot"]],
            ["Before plot values", dataset["before_bins"]],
            ["After plot values", dataset["after_bins"]],
            ["Before seed/fold values", dataset["before_raw"]],
            ["After seed/fold values", dataset["after_raw"]],
        ],
        columns=["field", "value"],
    )

    row = 0
    pd.DataFrame([["Source files"]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=row
    )
    row += 1
    metadata.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
    metadata_header_row = row + 1
    row += len(metadata) + 3

    pd.DataFrame([["Plot values used by grouped bin PDFs"]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=row
    )
    plot_title_row = row + 1
    row += 1
    plot_values.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
    plot_header_row = row + 1
    row += len(plot_values) + 3

    pd.DataFrame(
        [["Seed/fold summary: mean/std across seeds after averaging folds per seed"]]
    ).to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=row)
    seed_summary_title_row = row + 1
    row += 1
    seed_summary.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
    seed_summary_header_row = row + 1
    row += len(seed_summary) + 3

    pd.DataFrame([["Raw seed/fold values"]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=row
    )
    raw_title_row = row + 1
    row += 1
    raw_values.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row)
    raw_header_row = row + 1

    worksheet = writer.sheets[sheet_name]
    style_sheet(
        worksheet,
        title_rows=[1, plot_title_row, seed_summary_title_row, raw_title_row],
        header_rows=[metadata_header_row, plot_header_row, seed_summary_header_row, raw_header_row],
    )


def export_workbook(results_dir: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for dataset in DATASETS:
            plot_values = load_plot_values(results_dir, dataset)
            raw_values = load_seed_fold_values(results_dir, dataset)
            seed_summary = summarize_seed_fold(raw_values)
            write_dataset_sheet(
                writer=writer,
                sheet_name=dataset["sheet"][:MAX_SHEET_NAME_LEN],
                plot_values=plot_values,
                seed_summary=seed_summary,
                raw_values=raw_values,
                dataset=dataset,
            )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create an Excel workbook with R1-q8 grouped bin plot values."
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=default_results_dir(),
        help="R1-q8 Results folder.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to <results-dir>/grouped_bin_plot_values.xlsx.",
    )
    args = parser.parse_args()

    output_path = args.output or args.results_dir / "grouped_bin_plot_values.xlsx"
    export_workbook(args.results_dir, output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
