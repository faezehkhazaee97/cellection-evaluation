#!/usr/bin/env python3
"""Collect the numbers behind every Figure S7 ablation plot into one Excel workbook.

Run it once all five training scripts have finished: for each dataset in DATASETS (Cardio,
Cardio_Biosample, Covid, Lupus_Disease, Lupus_Race) it reads ablation_results.csv and
ablation_pivot_{AUROC,AUPRC,F1}.csv from --results-dir (default $RESULTS_DIR/R2-Q4/Results).
The output, <results-dir>/ablation_plot_values.xlsx or --output, has one sheet per dataset with
five stacked tables: source files and plot names; curve values (mean, std and row count of
AUROC, AUPRC and F1 per setting x fraction); delta bar-plot values (mean and std of top, random
and delta per fraction); the raw ablation rows; and the raw pivot rows. Needs pandas and
openpyxl, no GPU.
"""

# Paths come from environment variables (see experiments-public/README.md).
import os
RESULTS_DIR = os.environ.get("RESULTS_DIR", "results")  # output root

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


DATASETS = [
    ("Cardio", "Cardio", "Cardio"),
    ("Cardio_Biosample", "Cardio Biosample", "CardioBio"),
    ("Covid", "COVID", "Covid"),
    ("Lupus_Disease", "Lupus Disease", "LupusDis"),
    ("Lupus_Race", "Lupus Race", "LupusRace"),
]

CURVE_PLOTS = [
    ("AUROC", "AUROC", "ablation_auroc_curve.pdf", "AUROC Curve"),
    ("AUPRC", "AUPRC", "ablation_auprc_curve.pdf", "AUPRC Curve"),
    ("F1", "MacroF1", "ablation_f1_curve.pdf", "MacroF1 Curve"),
]

DELTA_PLOTS = [
    ("AUROC", "AUROC", "ablation_delta_auroc_barplot.pdf", "Delta AUROC"),
    ("AUPRC", "AUPRC", "ablation_delta_auprc_barplot.pdf", "Delta AUPRC"),
    ("F1", "MacroF1", "ablation_delta_f1_barplot.pdf", "Delta F1"),
]

# Excel limits sheet names to 31 characters.
EXCEL_SHEET_NAME_MAX_LEN = 31
# Upper bound on the auto-sized column width (in characters).
MAX_COLUMN_WIDTH = 40


def default_results_dir() -> Path:
    return Path(RESULTS_DIR) / "R2-Q4" / "Results"


def summarize_curve(df_ab: pd.DataFrame, metric: str) -> pd.DataFrame:
    """Mean, std and row count of `metric` per setting x frac_removed (the curve plot values)."""
    grouped = df_ab.groupby(["setting", "frac_removed"], sort=True)[metric]
    summary = grouped.agg(["mean", "std", "count"]).reset_index()
    summary = summary.rename(
        columns={
            "mean": f"{metric}_mean",
            "std": f"{metric}_std",
            "count": "n_rows",
        }
    )
    return summary


def raw_curve_values(df_ab: pd.DataFrame, metric: str) -> pd.DataFrame:
    """Rows behind one curve plot: seed, fold, frac_removed, setting, `metric` and any baseline columns."""
    columns = [
        "seed",
        "fold",
        "frac_removed",
        "setting",
        metric,
    ]
    for baseline_col in [f"baseline_{metric}", "baseline_AUROC", "baseline_AUPRC"]:
        if baseline_col in df_ab.columns and baseline_col not in columns:
            columns.append(baseline_col)
    columns = [column for column in columns if column in df_ab.columns]
    return df_ab[columns].sort_values(
        ["setting", "frac_removed", "seed", "fold"], kind="stable"
    )


def summarize_delta(pivot_df: pd.DataFrame) -> pd.DataFrame:
    """Mean and std of top, random and delta per frac_removed (the delta bar-plot values)."""
    grouped = pivot_df.groupby("frac_removed", sort=True)
    summary = grouped.agg(
        top_mean=("top", "mean"),
        top_std=("top", "std"),
        random_mean=("random", "mean"),
        random_std=("random", "std"),
        delta_mean=("delta", "mean"),
        delta_std=("delta", "std"),
        n_rows=("delta", "count"),
    ).reset_index()
    return summary


def all_curve_summaries(df_ab: pd.DataFrame) -> pd.DataFrame:
    frames = []
    for metric, metric_label, plot_file, _ in CURVE_PLOTS:
        summary = summarize_curve(df_ab, metric)
        summary.insert(0, "metric", metric_label)
        summary.insert(1, "plot_file", plot_file)
        summary = summary.rename(
            columns={
                f"{metric}_mean": "value_mean",
                f"{metric}_std": "value_std",
            }
        )
        frames.append(summary)
    return pd.concat(frames, ignore_index=True)


def all_delta_summaries(dataset_dir: Path) -> pd.DataFrame:
    frames = []
    for metric, metric_label, plot_file, _ in DELTA_PLOTS:
        pivot_df = pd.read_csv(dataset_dir / f"ablation_pivot_{metric}.csv")
        summary = summarize_delta(pivot_df)
        summary.insert(0, "metric", metric_label)
        summary.insert(1, "plot_file", plot_file)
        frames.append(summary)
    return pd.concat(frames, ignore_index=True)


def all_delta_raw(dataset_dir: Path) -> pd.DataFrame:
    frames = []
    for metric, metric_label, plot_file, _ in DELTA_PLOTS:
        pivot_df = pd.read_csv(dataset_dir / f"ablation_pivot_{metric}.csv")
        pivot_df.insert(0, "metric", metric_label)
        pivot_df.insert(1, "plot_file", plot_file)
        frames.append(pivot_df)
    return pd.concat(frames, ignore_index=True).sort_values(
        ["metric", "frac_removed", "seed", "fold"], kind="stable"
    )


def plot_files_table(dataset_label: str) -> pd.DataFrame:
    rows = [["Dataset", dataset_label]]
    for _, metric_label, plot_file, _ in CURVE_PLOTS:
        rows.append([f"{metric_label} curve plot", plot_file])
    for _, metric_label, plot_file, _ in DELTA_PLOTS:
        rows.append([f"Delta {metric_label} barplot", plot_file])
    rows.extend(
        [
            ["Curve source CSV", "ablation_results.csv"],
            ["Delta source CSVs", "ablation_pivot_AUROC.csv; ablation_pivot_AUPRC.csv; ablation_pivot_F1.csv"],
        ]
    )
    return pd.DataFrame(rows, columns=["field", "value"])


def autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, MAX_COLUMN_WIDTH)


def style_sheet(worksheet, title_rows: list[int], header_rows: list[int]) -> None:
    title_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_fill = PatternFill(fill_type="solid", fgColor="EDEDED")

    for row_idx in title_rows:
        for cell in worksheet[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = title_fill

    for row_idx in header_rows:
        for cell in worksheet[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

    worksheet.freeze_panes = "A3"
    autosize_columns(worksheet)


def write_table(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    table: pd.DataFrame,
    startrow: int,
) -> tuple[int, int, int]:
    """Write a title row and `table` below it, starting at 0-based row `startrow`.

    Returns (next_start, title_row, header_row): the 0-based row for the next table (two blank
    rows in between) and the 1-based Excel rows of the title and the header, for styling.
    """
    pd.DataFrame([[title]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=startrow
    )
    table.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow + 1)
    title_row = startrow + 1
    header_row = startrow + 2
    next_start = startrow + len(table) + 4
    return next_start, title_row, header_row


def write_dataset_sheet(
    writer: pd.ExcelWriter,
    sheet_name: str,
    dataset_dir: Path,
    dataset_label: str,
    df_ab: pd.DataFrame,
) -> None:
    pd.DataFrame([[f"{dataset_label} - Ablation Plot Values"]]).to_excel(
        writer, sheet_name=sheet_name, index=False, header=False, startrow=0
    )
    row = 2
    title_rows = [1]
    header_rows = []

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Source files and plot definitions",
        plot_files_table(dataset_label),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Curve plot values for AUROC, AUPRC, and MacroF1",
        all_curve_summaries(df_ab),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Delta barplot values for AUROC, AUPRC, and MacroF1",
        all_delta_summaries(dataset_dir),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Raw ablation rows used for curve plots",
        df_ab.sort_values(["setting", "frac_removed", "seed", "fold"], kind="stable"),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    row, title_row, header_row = write_table(
        writer,
        sheet_name,
        "Raw pivot rows used for delta barplots",
        all_delta_raw(dataset_dir),
        row,
    )
    title_rows.append(title_row)
    header_rows.append(header_row)

    style_sheet(writer.sheets[sheet_name], title_rows, header_rows)


def export_workbook(results_dir: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for dataset_dir_name, dataset_label, sheet_prefix in DATASETS:
            dataset_dir = results_dir / dataset_dir_name
            df_ab = pd.read_csv(dataset_dir / "ablation_results.csv")
            write_dataset_sheet(
                writer=writer,
                sheet_name=sheet_prefix[:EXCEL_SHEET_NAME_MAX_LEN],
                dataset_dir=dataset_dir,
                dataset_label=dataset_label,
                df_ab=df_ab,
            )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create an Excel workbook for R2-Q4 ablation plot values."
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=default_results_dir(),
        help="R2-Q4 Results folder.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to <results-dir>/ablation_plot_values.xlsx.",
    )
    args = parser.parse_args()

    output_path = args.output or args.results_dir / "ablation_plot_values.xlsx"
    export_workbook(args.results_dir, output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
